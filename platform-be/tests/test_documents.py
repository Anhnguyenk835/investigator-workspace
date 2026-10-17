"""Attached documents become evidence; a file that cannot be read is skipped, never guessed."""

from __future__ import annotations

import base64
import io
from typing import Any

import httpx
import pytest
from docx import Document

from investigator.connectors import documents
from investigator.connectors.documents import extract_text, supported
from investigator.connectors.gmail import GmailConnector
from investigator.connectors.support import SourceHttpError, search_terms
from investigator.tools import render_record
from tests.fakes import rec
from tests.test_connectors import CTX, Net, Tokens, gmail_message


def tiny_pdf(text: str) -> bytes:
    """A one-page PDF with `text` on it, with a correct index."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
    objects = [
        "<</Type/Catalog/Pages 2 0 R>>",
        "<</Type/Pages/Kids[3 0 R]/Count 1>>",
        (
            "<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
            "/Resources<</Font<</F1 5 0 R>>>>>>"
        ),
        f"<</Length {len(stream)}>>\nstream\n{stream}\nendstream",
        "<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<</Root 1 0 R/Size {len(objects) + 1}>>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def docx_bytes(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---- extraction -----------------------------------------------------------------------


def test_pdf_docx_and_plain_text_are_read() -> None:
    assert "Subscription form v3" in (
        extract_text("application/pdf", "f.pdf", tiny_pdf("Subscription form v3")) or ""
    )
    doc = extract_text(
        "",
        "terms.docx",
        docx_bytes(
            "Side letter",
            "MFN applies",
            table=[["Investor", "Amount"], ["Oakline", "5M"]],
        ),
    )
    assert (
        doc and "Side letter" in doc and "MFN applies" in doc and "Oakline | 5M" in doc
    )
    assert extract_text("text/plain", "n.txt", b"a  b\n\n\n\nc") == "a b\n\nc"
    assert (
        extract_text("text/html", "p.html", b"<p>Hello <b>there</b></p>")
        == "Hello there"
    )


def test_what_cannot_be_read_is_skipped_not_guessed() -> None:
    assert extract_text("image/png", "scan.png", b"\x89PNG") is None
    assert extract_text("application/pdf", "bad.pdf", b"not a pdf at all") is None
    assert extract_text("", "bad.docx", b"not a zip") is None
    assert (
        extract_text("text/plain", "big.txt", b"x" * (documents.MAX_BYTES + 1)) is None
    )
    assert extract_text("text/plain", "empty.txt", b"   \n ") is None
    assert not supported("image/jpeg", "photo.jpg") and supported("", "FORM.PDF")


def test_very_long_text_is_cut() -> None:
    text = extract_text("text/plain", "long.txt", b"word " * 50_000)
    assert text is not None and len(text) <= documents.MAX_CHARS


# ---- through Gmail --------------------------------------------------------------------


def with_attachments(*files: tuple[str, str, str]) -> dict[str, Any]:
    """A message carrying (partId, filename, mimeType) attachments."""
    msg = gmail_message("m1")
    for part_id, name, mime in files:
        msg["payload"]["parts"].append(
            {
                "partId": part_id,
                "mimeType": mime,
                "filename": name,
                "body": {"attachmentId": f"att-{part_id}", "size": 100},
            }
        )
    return msg


def attachment_route(blobs: dict[str, bytes]) -> Any:
    def route(request: httpx.Request) -> httpx.Response:
        att = str(request.url).rsplit("/", 1)[-1].split("?")[0]
        data = base64.urlsafe_b64encode(blobs[att]).decode()
        return httpx.Response(200, json={"data": data})

    return route


async def test_a_thread_includes_its_readable_attachments_as_records() -> None:
    msg = with_attachments(
        ("2", "form.docx", "application/octet-stream"), ("3", "scan.png", "image/png")
    )
    net = Net(
        **{
            "/threads/t1": {"messages": [msg]},
            "/attachments/att-2": attachment_route(
                {"att-2": docx_bytes("Investor must re-sign v3")}
            ),
        }
    )
    thread = await GmailConnector(Tokens(), net.client).get_thread("gmail:t1", CTX)
    ids = [r.id for r in thread.records]
    assert ids == ["gmail:m1", "gmail:m1:att:2"] or set(ids) == {
        "gmail:m1",
        "gmail:m1:att:2",
    }
    doc = next(r for r in thread.records if r.id.endswith(":att:2"))
    assert doc.title == "form.docx" and "Investor must re-sign v3" in doc.excerpt
    assert doc.thread_id == "gmail:t1" and doc.meta["kind"] == "attachment"
    assert doc.author == "Maria Chen" and doc.occurred_at is not None
    assert not any(
        "att-3" in str(r.url) for r in net.requests
    )  # an image is never downloaded
    assert thread.coverage.status == "searched"


async def test_only_a_few_attachments_are_read_per_thread_and_the_rest_are_reported() -> (
    None
):
    files = [(str(i), f"f{i}.txt", "text/plain") for i in range(2, 7)]
    msg = with_attachments(*files)
    blobs = {f"att-{i}": f"text {i}".encode() for i in range(2, 7)}
    net = Net(
        **{"/threads/t1": {"messages": [msg]}, "/attachments/": attachment_route(blobs)}
    )
    thread = await GmailConnector(Tokens(), net.client).get_thread("gmail:t1", CTX)
    assert len([r for r in thread.records if ":att:" in r.id]) == 3
    assert thread.coverage.status == "partial" and "2 attachment(s) were not read" in (
        thread.coverage.detail or ""
    )


async def test_one_attachment_by_id_and_bad_ids() -> None:
    msg = with_attachments(("2", "form.pdf", "application/pdf"))
    net = Net(
        **{
            "/attachments/att-2": attachment_route(
                {"att-2": tiny_pdf("Wire details changed")}
            ),
            "/messages/m1": msg,
        }
    )
    gmail = GmailConnector(Tokens(), net.client)
    got = await gmail.get_record("gmail:m1:att:2", CTX)
    assert "Wire details changed" in got.records[0].excerpt
    assert (await gmail.get_record("gmail:m1:att:9", CTX)).records == []
    assert (await gmail.get_record("gmail:m1:att:../x", CTX)).records == []


async def test_an_unreadable_attachment_is_reported_not_invented() -> None:
    msg = with_attachments(("2", "bad.pdf", "application/pdf"))
    net = Net(
        **{
            "/attachments/att-2": attachment_route({"att-2": b"garbage"}),
            "/messages/m1": msg,
        }
    )
    got = await GmailConnector(Tokens(), net.client).get_record("gmail:m1:att:2", CTX)
    assert got.records == [] and got.coverage.status == "partial"


def test_documents_get_more_room_than_messages_when_shown_to_the_model() -> None:
    long_text = "x" * 3000
    plain = render_record(rec("m", excerpt=long_text), 1500)
    doc = render_record(rec("d", excerpt=long_text, meta={"kind": "attachment"}), 1500)
    assert "…[truncated]" in plain and "…[truncated]" not in doc


# ---- other recall helpers -------------------------------------------------------------


def test_vietnamese_filler_words_are_dropped_from_queries() -> None:
    assert search_terms("tại sao biểu mẫu bị trễ") == ["biểu", "mẫu", "trễ"]
    assert search_terms("của và là") == ["của", "và", "là"]  # never empty


async def test_a_short_rate_limit_is_waited_out_once_and_a_long_one_is_not(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    waits: list[float] = []

    async def fake_sleep(seconds: float) -> None:
        waits.append(seconds)

    monkeypatch.setattr("investigator.connectors.support.asyncio.sleep", fake_sleep)
    calls = {"n": 0}

    def limited(request: httpx.Request) -> httpx.Response:
        calls["n"] += 1
        if calls["n"] == 1:
            return httpx.Response(429, headers={"retry-after": "2"}, json={})
        return httpx.Response(200, json={"labels": []})

    net = Net(**{"/labels": limited})
    gmail = GmailConnector(Tokens(), net.client)
    status, _ = await gmail.suggestions(CTX)
    assert status == "ok" and waits == [2.0] and calls["n"] == 2

    waits.clear()
    long_wait = Net(
        **{
            "/labels": lambda r: httpx.Response(
                429, headers={"retry-after": "60"}, json={}
            )
        }
    )
    status, _ = await GmailConnector(Tokens(), long_wait.client).suggestions(CTX)
    assert status == "failed" and waits == []

    waits.clear()
    always = Net(**{"/labels": lambda r: httpx.Response(503, json={})})
    status, _ = await GmailConnector(Tokens(), always.client).suggestions(CTX)
    assert status == "failed" and len(waits) == 2  # retried twice, then gave up

    with pytest.raises(SourceHttpError):
        await GmailConnector(Tokens(), always.client)._get(
            "https://x.test/y", "tok-secret"
        )
