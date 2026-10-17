"""Small real documents for the mail: a PDF, a Word file, a CSV. Built in memory."""

from __future__ import annotations

import io
import textwrap

from investigator.devseed.beats import Attach


def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf(title: str, lines: list[str]) -> bytes:
    """A minimal text PDF (Helvetica, A4, several pages): enough for a text extractor."""
    wrapped: list[str] = []
    for ln in lines:
        wrapped += textwrap.wrap(ln, 92) or [""]
    pages = [wrapped[i : i + 46] for i in range(0, len(wrapped), 46)] or [[]]
    objs: list[bytes] = []

    def add(body: bytes) -> int:
        objs.append(body)
        return len(objs)

    add(b"")  # 1: catalog (filled later)
    add(b"")  # 2: pages (filled later)
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids: list[int] = []
    for n, page in enumerate(pages):
        stream = ["BT", "/F1 11 Tf", "50 790 Td", "14 TL"]
        if n == 0 and title:
            stream += ["/F1 15 Tf", f"({_esc(title)}) Tj", "T*", "/F1 11 Tf", "T*"]
        stream += [f"({_esc(ln)}) Tj T*" for ln in page]
        stream.append("ET")
        data = "\n".join(stream).encode("latin-1", "replace")
        content = add(b"<< /Length %d >>\nstream\n" % len(data) + data + b"\nendstream")
        page_ids.append(
            add(
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {content} 0 R "
                f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
            )
        )
    objs[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objs[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def docx(title: str, lines: list[str]) -> bytes:
    from docx import Document

    d = Document()
    if title:
        d.add_heading(title, level=1)
    for ln in lines:
        d.add_paragraph(ln)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build(a: Attach) -> tuple[bytes, str, str]:
    """(bytes, maintype, subtype)"""
    if a.kind == "pdf":
        return pdf(a.title, a.lines), "application", "pdf"
    if a.kind == "docx":
        return (
            docx(a.title, a.lines),
            "application",
            "vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    if a.kind == "csv":
        return ("\n".join(a.lines) + "\n").encode(), "text", "csv"
    return ("\n".join([a.title, *a.lines])).encode(), "text", "plain"
