"""Plain text from the documents people attach to mail. Untrusted input: strict size and
page limits, and a file that cannot be read is skipped rather than guessed at."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable

from investigator.connectors.support import html_to_text

MAX_BYTES = 6_000_000
MAX_CHARS = 20_000
MAX_PDF_PAGES = 40
MAX_DOCX_PARAGRAPHS = 2000

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_TYPES = {
    "text/plain",
    "text/csv",
    "text/markdown",
    "text/html",
    "application/json",
}
SUPPORTED = {PDF, DOCX, *TEXT_TYPES}


def supported(mime: str, filename: str) -> bool:
    return _kind(mime, filename) is not None


def _kind(mime: str, filename: str) -> str | None:
    mime = mime.lower().split(";")[0].strip()
    name = filename.lower()
    if mime == PDF or name.endswith(".pdf"):
        return "pdf"
    if mime == DOCX or name.endswith(".docx"):
        return "docx"
    if mime in TEXT_TYPES or name.endswith((".txt", ".csv", ".md", ".html", ".htm")):
        return "html" if "html" in mime or name.endswith((".html", ".htm")) else "text"
    return None


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        return ""
    pages = reader.pages[:MAX_PDF_PAGES]
    return "\n\n".join((p.extract_text() or "") for p in pages)


def _docx(data: bytes) -> str:
    from docx import Document

    # A .docx is a zip; refuse one that claims to expand absurdly (a zip bomb).
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist()) > 10 * MAX_BYTES:
            return ""
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs[:MAX_DOCX_PARAGRAPHS]]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _text(data: bytes) -> str:
    return data.decode("utf-8", "replace")


READERS: dict[str, Callable[[bytes], str]] = {
    "pdf": _pdf,
    "docx": _docx,
    "text": _text,
    "html": lambda d: html_to_text(_text(d)),
}


def extract_text(mime: str, filename: str, data: bytes) -> str | None:
    """The document's text, or None when it is too large, unsupported or unreadable."""
    kind = _kind(mime, filename)
    if kind is None or len(data) > MAX_BYTES:
        return None
    try:
        text = READERS[kind](data)
    except Exception:  # noqa: BLE001  a malformed file must not end the investigation
        return None
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:MAX_CHARS] if text else None
